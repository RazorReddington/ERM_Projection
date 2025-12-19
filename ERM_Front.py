# -*- coding: utf-8 -*-
"""
Created on Wed Jul 30 10:51:04 2025
@author: UG423NJ
"""

import streamlit as st
import pandas as pd
import numpy as np
import os
import subprocess
from openpyxl import load_workbook
import plotly.graph_objects as go
import time
from datetime import datetime


# --- Configuration ---
st.set_page_config(page_title="ERM Projection Interface", layout="wide")


# --- Load Scenario Parameters ---
#@st.cache_data
def load_scenarios(file_path, sheet_name):
    df = pd.read_excel(file_path, sheet_name, header=0)
    df.columns = [str(col).strip() for col in df.columns]
    # Rename duplicate columns
    seen = {}
    new_columns = []
    for col in df.columns:
        if col in seen:
            seen[col] += 1
            new_columns.append(f"{col}_{seen[col]}")
        else:
            seen[col] = 0
            new_columns.append(col)
    df.columns = new_columns
    return df



@st.cache_data
def load_mpf_summary(file_path):
    mpf = pd.read_csv(file_path)
   
    
    count_model_points = len(mpf)
    count_joint_life = sum(mpf['Joint Life'] == 'Joint Life')
    count_single_life = sum(mpf['Joint Life'] == 'Single')
    count_male1 = sum(mpf['Gender 1'] == "Male")
    count_female1 = sum(mpf['Gender 1'] == "Female")
    count_male2 = sum(mpf['Gender 2'] == "Male")
    count_female2 = sum(mpf['Gender 2'] == "Female")
    average_age1 = np.average(mpf['Age 1'])
    average_age2 = np.nanmean(mpf['Age 2'])
    weighted_ltv = sum(mpf['LTV'] * mpf['Loan Amount'])/sum(mpf['Loan Amount'])
    weighted_aer = sum(mpf['AER'] * mpf['Loan Amount'])/sum(mpf['Loan Amount'])
    average_property = np.average(mpf['Loan Amount'] / mpf['LTV'])
    
    mpf_summary = pd.DataFrame({
    "Metric": [
        "Total Model Points",
        "Joint Life Count",
        "Single Life Count",
        "Male (Primary)",
        "Female (Primary)",
        "Male (Secondary)",
        "Female (Secondary)",
        "Average Age (Primary)",
        "Average Age (Secondary",
        "Weighted LTV",
        "Weighted AER",
        "Average Property Value"
    ],
    "Value": [
        count_model_points,
        count_joint_life,
        count_single_life,
        count_male1,
        count_female1,
        count_male2,
        count_female2,
        average_age1,
        average_age2,
        weighted_ltv,
        weighted_aer,
        average_property
    ]
})
    return mpf_summary


# --- Load Model Output ---
#@st.cache_data
def load_output(file_path):
    cashflows = pd.read_excel(file_path, sheet_name="Cashflows")
    olb = pd.read_excel(file_path, sheet_name="OLB")
    return cashflows, olb



def save_edited_parameters(df, sheet_name, file_path):
    book = load_workbook(file_path)

    if sheet_name in book.sheetnames:
        std = book[sheet_name]
        book.remove(std)

    temp_path = file_path.replace(".xlsx", "_temp.xlsx")
    book.save(temp_path)

    with pd.ExcelWriter(temp_path, engine='openpyxl', mode='a') as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)

    os.replace(temp_path, file_path)
    st.success(f"Edited parameters saved to sheet '{sheet_name}' in {file_path}")




# --- Run ERM Model ---
def run_model(filepath):
    result = subprocess.run(["python", filepath], capture_output=True, text=True)
    if result.returncode == 0:
        st.success("Model executed successfully.")
        return load_output(Output)
        
 
    
    else:
        st.error("Model execution failed.")
        st.text(result.stderr)
        return None, None



########################## Dashboard Design #################################

st.title("ERM Projection")
page = st.sidebar.radio("Go to", ["File Paths", "Global Parameters", "Scenario Parameters", "Model Point File", "Model Output"])


if page == "File Paths":
    st.write("Set File Paths")
    variable_names = ["Master Input Filepath", "MPF Filepath", "Output Filepath", "Model Filepath"]

    for var in variable_names:
        st.session_state[var] = st.text_input(f"Enter path for {var}", st.session_state.get(var, ""))

    st.write("Collected File Paths:")
    st.write({var: st.session_state.get(var) for var in variable_names})

Master_Input = st.session_state.get("Master Input Filepath", "")
MPF = st.session_state.get("MPF Filepath", "")
Output = st.session_state.get("Output Filepath", "")
model_filepath = st.session_state.get("Model Filepath", "")

global_df = load_scenarios(Master_Input, 'Global_Parameters') #this raises an error on the dashboard before loading the filepaths - should look to resolve 
#model_filepath = global_df.loc[global_df['Parameter'] == "Model Filepath", 'Value'].iloc[0]
working_directory = global_df.loc[global_df['Parameter'] == "Working Directory", 'Value'].iloc[0]


if page == "Global Parameters":
    st.write("Set Global Parameters")
    global_df = load_scenarios(Master_Input, 'Global_Parameters')

    
    global_df[global_df.columns[1]] = global_df[global_df.columns[1]].astype(str) #format the value column as string to allow editing. Model script converts to necessary formats
    edited_global_df = st.data_editor(
        global_df,
        num_rows="dynamic",
        use_container_width=True,
        hide_index=True,
        key="global_editor",
        column_config={
                    global_df.columns[0]: st.column_config.TextColumn(disabled=True),
                    global_df.columns[1]: st.column_config.TextColumn(disabled=False)
                }

    )

    st.session_state["global_params"] = edited_global_df
    
    if st.button("Save Global Parameters"):
         save_edited_parameters(
             edited_global_df,
             sheet_name="Global_Parameters",
             file_path= Master_Input
         )


if page == "Scenario Parameters":
    st.write("Set Scenario Parameters")
    scenario_df = load_scenarios(Master_Input, 'Scenario_Parameters')
    edited_df = st.data_editor(
        scenario_df,
        num_rows="dynamic",
        use_container_width=True,
        hide_index=True,
        key="scenario_editor"
    )

    
    st.session_state["edited_scenarios"] = edited_df
    
    if st.button("Save Scenario Parameters"):
         save_edited_parameters(
             edited_df,
             sheet_name="Scenario_Parameters",
             file_path= Master_Input
         )
         

         
if page == "Model Point File":
    st.write("Review Model Points")
    model_point_summary = load_mpf_summary(MPF) 
    st.session_state["MPF"] = model_point_summary
    st.dataframe(model_point_summary, hide_index=True)
    
    
if page == "Model Output":
    st.subheader("Run Model and View Output")
    if st.button("Run ERM Model"):
        username = os.environ.get('USERNAME')
        runtime = datetime.now()
        
        
      # Save to audit log file
        log_entry = f"{runtime} | User: {username}\n"
        file_path = os.path.join(working_directory, "audit_log.txt")
        with open(file_path, "a") as log_file:  # Append mode
            log_file.write(log_entry)

        st.success(f"Model run logged for {username} at {runtime}")

        
        st.session_state["cashflows_df"] = None
        st.session_state["olb_df"] = None

        cashflows_df, olb_df = run_model(model_filepath)
        st.session_state["cashflows_df"] = cashflows_df
        st.session_state["olb_df"] = olb_df
        

    if "cashflows_df" in st.session_state and st.session_state["cashflows_df"] is not None and \
       "olb_df" in st.session_state and st.session_state["olb_df"] is not None:
    
        cashflows_df = st.session_state["cashflows_df"]
        olb_df = st.session_state["olb_df"]
    
        olb_column = st.selectbox("Select OLB column to plot", options=olb_df.columns[1:])
        cashflow_column = st.selectbox("Select Cashflow column to plot", options=cashflows_df.columns[1:])
    
        fig = go.Figure()
    
        fig.add_trace(go.Scatter(
            x=cashflows_df.iloc[:, 0],  # Time axis
            y=cashflows_df[cashflow_column],  # Cashflows
            name=f"{cashflow_column}",
            yaxis="y1"
        ))
    
        fig.add_trace(go.Scatter(
            x=olb_df.iloc[:, 0],  # Time axis
            y=olb_df[olb_column],  # OLB
            name=f"{olb_column}",
            yaxis="y2"
        ))
    
        fig.update_layout(
            title="Cashflows and OLB Projection",
            xaxis=dict(title="Time"),
            yaxis=dict(title="Cashflows", side="left"),
            yaxis2=dict(title="OLB", overlaying="y", side="right"),
            legend=dict(x=0.9, y=0.99)
        )
    
        st.plotly_chart(fig, use_container_width=True)



