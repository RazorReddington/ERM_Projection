# -*- coding: utf-8 -*-
"""
Created on Wed Jul 30 10:51:04 2025
@author: UG423NJ
"""

import streamlit as st
import pandas as pd
import numpy as np
import os
import subprocess
from openpyxl import load_workbook


# --- Configuration ---
st.set_page_config(page_title="ERM Projection Interface", layout="wide")

# --- Load Scenario Parameters ---
#@st.cache_data
def load_scenarios(file_path, sheet_name):
    df = pd.read_excel(file_path, sheet_name, header=0)
    df.columns = [str(col).strip() for col in df.columns]
    # Rename duplicate columns
    seen = {}
    new_columns = []
    for col in df.columns:
        if col in seen:
            seen[col] += 1
            new_columns.append(f"{col}_{seen[col]}")
        else:
            seen[col] = 0
            new_columns.append(col)
    df.columns = new_columns
    return df





@st.cache_data
def load_mpf_summary(file_path):
    mpf = pd.read_csv(file_path)
   
    
    count_model_points = len(mpf)
    count_joint_life = sum(mpf['Joint Life'] == 'Joint Life')
    count_single_life = sum(mpf['Joint Life'] == 'Single')
    count_male1 = sum(mpf['Gender 1'] == "Male")
    count_female1 = sum(mpf['Gender 1'] == "Female")
    count_male2 = sum(mpf['Gender 2'] == "Male")
    count_female2 = sum(mpf['Gender 2'] == "Female")
    average_age1 = np.average(mpf['Age 1'])
    average_age2 = np.nanmean(mpf['Age 2'])
    weighted_ltv = sum(mpf['LTV'] * mpf['Loan Amount'])/sum(mpf['Loan Amount'])
    weighted_aer = sum(mpf['AER'] * mpf['Loan Amount'])/sum(mpf['Loan Amount'])
    average_property = np.average(mpf['Loan Amount'] / mpf['LTV'])
    
    mpf_summary = pd.DataFrame({
    "Metric": [
        "Total Model Points",
        "Joint Life Count",
        "Single Life Count",
        "Male (Primary)",
        "Female (Primary)",
        "Male (Secondary)",
        "Female (Secondary)",
        "Average Age (Primary)",
        "Average Age (Secondary",
        "Weighted LTV",
        "Weighted AER",
        "Average Property Value"
    ],
    "Value": [
        count_model_points,
        count_joint_life,
        count_single_life,
        count_male1,
        count_female1,
        count_male2,
        count_female2,
        average_age1,
        average_age2,
        weighted_ltv,
        weighted_aer,
        average_property
    ]
})
    return mpf_summary


# --- Load Model Output ---
@st.cache_data
def load_output(file_path):
    cashflows = pd.read_excel(file_path, sheet_name="Cashflows")
    olb = pd.read_excel(file_path, sheet_name="OLB")
    return cashflows, olb



def save_edited_parameters(df, sheet_name, file_path = 'C:/Users/UG423NJ/OneDrive - EY/Documents/GitHub/ERM_Projection/Data/Master_Input.xlsx'):
    # Load the existing workbook
    book = load_workbook(file_path)

    # Remove the sheet if it already exists to avoid duplication
    if sheet_name in book.sheetnames:
        std = book[sheet_name]
        book.remove(std)

    # Save the workbook temporarily without the target sheet
    temp_path = file_path.replace(".xlsx", "_temp.xlsx")
    book.save(temp_path)

    # Now write the updated sheet
    with pd.ExcelWriter(temp_path, engine='openpyxl', mode='a') as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)

    # Replace the original file with the updated one

    os.replace(temp_path, file_path)

    st.success(f"Edited parameters saved to sheet '{sheet_name}' in {file_path}")




# --- Run ERM Model ---
def run_model():
    result = subprocess.run(["python", "ERM_Projection.py"], capture_output=True, text=True)
    if result.returncode == 0:
        st.success("Model executed successfully.")
        return load_output("C:/Users/UG423NJ/OneDrive - EY/Documents/GitHub/ERM_Projection/Data/output.xlsx")
        
 
    
    else:
        st.error("Model execution failed.")
        st.text(result.stderr)
        return None, None




# --- UI Layout ---
st.title("ERM Projection")
page = st.sidebar.radio("Go to", ["Global Parameters", "Scenario Parameters", "Model Point File", "Model Output"])




if page == "Global Parameters":
    st.write("Set Global Parameters")
    
    global_df = load_scenarios("C:/Users/UG423NJ/OneDrive - EY/Documents/GitHub/ERM_Projection/Data/Master_Input.xlsx", 'Global_Parameters')
    global_df[global_df.columns[1]] = global_df[global_df.columns[1]].astype(str) #format the value column as string to allow editing. Model script converts to necessary formats
    edited_global_df = st.data_editor(
        global_df,
        num_rows="dynamic",
        use_container_width=True,
        hide_index=True,
        key="global_editor",
        column_config={
                    global_df.columns[0]: st.column_config.TextColumn(disabled=True),
                    global_df.columns[1]: st.column_config.TextColumn(disabled=False)
                }

    )

    st.session_state["global_params"] = edited_global_df
    
    if st.button("Save Global Parameters"):
         save_edited_parameters(
             edited_global_df,
             sheet_name="Global_Parameters",
             file_path="C:/Users/UG423NJ/OneDrive - EY/Documents/GitHub/ERM_Projection/Data/Master_Input.xlsx"
         )


if page == "Scenario Parameters":
    st.write("Set Scenario Parameters")
    scenario_df = load_scenarios("C:/Users/UG423NJ/OneDrive - EY/Documents/GitHub/ERM_Projection/Data/Master_Input.xlsx", 'Scenario_Parameters')
    edited_df = st.data_editor(
        scenario_df,
        num_rows="dynamic",
        use_container_width=True,
        hide_index=True,
        key="scenario_editor"
    )

    
    st.session_state["edited_scenarios"] = edited_df
    
    if st.button("Save Scenario Parameters"):
         save_edited_parameters(
             edited_df,
             sheet_name="Scenario_Parameters",
             file_path="C:/Users/UG423NJ/OneDrive - EY/Documents/GitHub/ERM_Projection/Data/Master_Input.xlsx"
         )
         
if page == "Model Point File":
    st.write("Review Model Points")
    model_point_summary = load_mpf_summary("C:/Users/UG423NJ/OneDrive - EY/Documents/GitHub/ERM_Projection/Data/MPF_Phoenix_Internal.csv") 
    st.session_state["MPF"] = model_point_summary
    st.dataframe(model_point_summary, hide_index=True)
    
    
if page == "Model Output":
    st.subheader("Run Model and View Output")
    if st.button("Run ERM Model"):
        cashflows_df, olb_df = run_model()


        if cashflows_df is not None:
            st.subheader("Cashflows")
            st.dataframe(cashflows_df)
            st.subheader("OLB")
            st.dataframe(olb_df)
            
   
    
    #selected_run = st.selectbox("Select a run to view chart", list(chart_dict.keys()))
    #st.pyplot(chart_dict[selected_run])



